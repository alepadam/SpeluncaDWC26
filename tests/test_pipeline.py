"""End-to-end test on a synthetic dataset that mirrors the real file structure (never the real data)."""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src.config import load_config
from src.pipeline import run_pipeline
from tests.synthetic import make_dataset

ROOT = Path(__file__).resolve().parents[1]
FAST = {"optimization": {"lhs_samples": 2000, "posterior_keep": 100, "de_maxiter": 25, "de_popsize": 10},
        "proxy": {"gp_restarts": 0}, "sensitivity": {"sobol_n": 16}}


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    root = tmp_path_factory.mktemp("spelunca")
    (root / "data" / "raw").mkdir(parents=True)
    make_dataset(root / "data" / "raw")
    cfg = load_config(ROOT / "config.yaml", root=root, overrides=FAST, use_env=False)
    return cfg, run_pipeline(cfg)


def test_convention_detected_and_history_column_excluded(run):
    cfg, st = run
    assert st["convention"] == "start"
    assert "RAMP_History" not in st["sim_cum"]["oil"].columns
    assert st["sim_cum"]["oil"].shape == (41, 100)


def test_proxy_is_accurate_on_validation(run):
    _, st = run
    s = st["proxy_scores"]
    r2 = s[(s.model == "pca_gp") & (s.split == "validation") & (s.phase == "oil")]["r2_end_2008"].iloc[0]
    assert r2 > 0.95


def test_optimizer_beats_best_simulated_case(run):
    _, st = run
    assert st["optimum"]["misfit_proxy"].iloc[0] <= st["case_misfit"]["misfit"].min()


def test_deliverable_is_complete_and_physical(run):
    cfg, st = run
    path = st["deliverable_path"]
    wb = load_workbook(path, read_only=True, data_only=True)
    assert wb.sheetnames == [f"Case {i}" for i in range(1, 11)]
    dates = pd.date_range("2008-04-01", "2028-01-01", freq="QS-JAN")
    obs_end = st["obs_cum"].iloc[-1]
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(min_row=2, max_row=81, max_col=4, values_only=True))
        assert [pd.Timestamp(r[0]) for r in rows] == list(dates)
        gas, oil, water = (np.array([r[i] for r in rows], float) for i in (1, 2, 3))
        for series, start in ((gas, obs_end["gas"]), (oil, obs_end["oil"]), (water, obs_end["water"])):
            assert np.all(np.diff(np.concatenate([[start], series])) >= 0)
    wb.close()


def test_all_result_tables_written(run):
    cfg, _ = run
    expected = ["srq1_convention_test.csv", "qc_report.csv", "srq4_proxy_scores.csv", "srq4_proxy_uncertainty.csv",
                "srq2_case_misfit.csv", "srq2_optimum.csv", "srq2_posterior_samples.csv", "srq2_sigma_sensitivity.csv",
                "srq5_backtest_summary.csv", "srq5_all_cases_2028.csv", "srq6_best10.csv", "srq6_spread_2028.csv",
                "srq6_business.csv", "srq6_pipeline_backtest.csv", "srq3_spearman.csv", "srq3_sobol.csv"]
    missing = [f for f in expected if not (cfg["paths"]["tables"] / f).is_file()]
    assert not missing


def test_matching_narrows_the_2028_forecast(run):
    _, st = run
    sp = st["spread"]
    oil = sp[sp.phase == "oil"].set_index("set")["P10_P90_width"]
    assert oil["posterior"] < oil["all_100_cases"]


def test_forecast_method_and_water_trend_selected_on_validation(run):
    cfg, st = run
    assert st["forecast_method"] in ("decline", "ml_autoregressive")
    assert st["water_model"] in ("wor_vs_np", "wor_vs_time")
    sel = pd.read_csv(cfg["paths"]["tables"] / "srq5_method_selection.csv")
    assert sel["selected"].sum() == 1
