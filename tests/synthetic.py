"""Synthetic DWC26-like dataset for tests and CI (never the real data).

Mirrors every structural fact confirmed from the real files:
  * 01 Introduction.xlsx: 'Nomenclature' + 'Uncertainty Values' (Item, 4 parameters, 100 cases)
  * 02 Field Production.xlsx: 'Field Production' (41 quarterly rows, start-of-quarter rate stamps)
    + curve sheets with 3 header rows (Y1..Y101 / Case_1..Case_100 + RAMP_History / units),
      irregular timestamps, blanks where a case did not report, RAMP_History at quarter dates only
  * 09 Template Deliverable.xlsx: sheets 'Case 1'..'Case 10', 80 quarters, text dates except the last
  * files placed under a sub-folder (like dwc26/TestCases/) to exercise the file search
Physics is simplified (plateau + hyperbolic decline, waterflood WOR growth, constant GOR 0.3633).
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

RS = 0.3633
Q_PLATEAU = 15507.77
START, END = pd.Timestamp("1998-01-01"), pd.Timestamp("2008-01-01")
PROD_START = pd.Timestamp("1998-10-01")
PARAMS = ["Fault Transmissibility", "Porosity Multiplier", "Permeability Multiplier", "Aquifer Pore Volume"]
BOUNDS = [(0.05, 0.15), (0.8, 1.5), (0.5, 9.8), (54.0, 200.0)]


def _case_daily(p: np.ndarray, days: pd.DatetimeIndex) -> dict[str, np.ndarray]:
    """Daily oil, gas, water rates for one parameter set."""
    flt, poro, perm, aq = p
    t = (days - PROD_START).days.to_numpy(float)
    on = t >= 0
    plateau = 900.0 * poro * (1.0 + 2.0 * flt) * (0.85 + 0.03 * np.log(perm))
    di = 0.00045 * (1.4 - 0.25 * np.log(perm)) / poro
    b = 0.5
    tt = np.clip(t - plateau, 0.0, None)
    qo = np.where(on, Q_PLATEAU / np.power(1.0 + b * di * tt, 1.0 / b), 0.0)
    tb = 700.0 - 60.0 * np.log(perm) - 1.2 * (aq - 54.0)
    wor = np.where(t > tb, 0.02 * np.exp((t - tb) / 365.0 * (0.45 + aq / 400.0)), 0.0)
    return {"oil": qo, "gas": RS * qo, "water": np.minimum(wor, 50.0) * qo}


def _cumulative_at(rates: np.ndarray, days: pd.DatetimeIndex, stamps: pd.DatetimeIndex) -> np.ndarray:
    """Volume produced from START up to each stamp (daily rate integration, linear within the day)."""
    cum = np.concatenate([[0.0], np.cumsum(rates)])        # cum[i] = volume up to days[0] + i days
    pos = (stamps - days[0]).total_seconds().to_numpy() / 86400.0
    return np.interp(pos, np.arange(cum.size), cum)


def _write_curve_sheet(ws, stamps, columns, units):
    ws.append(["X"] + [f"Y{i + 1}" for i in range(len(columns))])
    ws.append(["Date"] + list(columns))
    ws.append([None] + [units] * len(columns))
    data = np.column_stack([columns[c] for c in columns])
    for i, s in enumerate(stamps):
        ws.append([s.to_pydatetime()] + [None if np.isnan(v) else float(v) for v in data[i]])


def make_dataset(root: Path, n_cases: int = 100, seed: int = 0) -> dict:
    """Write the synthetic files under root/dwc26/TestCases and return the hidden truth."""
    rng = np.random.default_rng(seed)
    out = Path(root) / "dwc26" / "TestCases"
    out.mkdir(parents=True, exist_ok=True)

    u = rng.random((n_cases, 4))
    lo, hi = np.array([b[0] for b in BOUNDS]), np.array([b[1] for b in BOUNDS])
    params = lo + u * (hi - lo)
    params[:, 2] = np.exp(np.log(lo[2]) + u[:, 2] * (np.log(hi[2]) - np.log(lo[2])))
    truth = np.array([0.10, 1.10, 3.0, 120.0])

    # ---- 01 Introduction
    with pd.ExcelWriter(out / "01 Introduction.xlsx") as xw:
        pd.DataFrame({"Grid": ["I Index"], "Name": [None]}).to_excel(xw, sheet_name="Nomenclature", index=False)
        df = pd.DataFrame(params, columns=PARAMS)
        df.insert(0, "Item", [f"Case {i + 1}" for i in range(n_cases)])
        df.to_excel(xw, sheet_name="Uncertainty Values", index=False)

    # ---- simulation
    days = pd.date_range(START, pd.Timestamp("2008-04-01"), freq="D")
    quarters = pd.date_range(START, END, freq="QS-JAN")
    extra = START + pd.to_timedelta(np.sort(rng.uniform(0, (END - START).days, 1500)), unit="D")
    stamps = quarters.union(pd.DatetimeIndex(extra).round("min")).unique()
    sim = {k: {} for k in ("oil", "gas", "water")}
    for i in range(n_cases):
        r = _case_daily(params[i], days)
        own = quarters.union(stamps[rng.random(len(stamps)) < 0.03])
        mask = stamps.isin(own)
        for k in sim:
            col = np.full(len(stamps), np.nan)
            col[mask] = _cumulative_at(r[k], days, stamps[mask])
            sim[k][f"Case_{i + 1}"] = col
    r_true = _case_daily(truth, days)
    for k in sim:
        col = np.full(len(stamps), np.nan)
        q_mask = stamps.isin(quarters)
        col[q_mask] = _cumulative_at(r_true[k], days, stamps[q_mask])
        sim[k]["RAMP_History"] = col

    # ---- observed: rate stamped at t_k = average over [t_k, t_k+1)  ('start' convention), 0.5% noise
    q_ext = pd.date_range(START, pd.Timestamp("2008-04-01"), freq="QS-JAN")
    obs = {}
    for k in ("oil", "gas", "water"):
        c = _cumulative_at(r_true[k], days, q_ext)
        rate = np.diff(c) / q_ext.to_series().diff().dt.days.to_numpy(float)[1:]
        obs[k] = rate * (1.0 + 0.005 * rng.standard_normal(rate.size))
    field = pd.DataFrame({
        "Identifier": "Field", "Date": quarters,
        "Oil production rate (STB/d)": np.round(obs["oil"], 2),
        "Gas production rate (MSCF/d)": np.round(obs["gas"], 2),
        "Water production rate (STB/d)": np.round(np.clip(obs["water"], 0, None), 2),
        "Water injection rate (STB/d)": np.where(quarters >= PROD_START, 14748.87, 0.0),
    })

    wb = Workbook()
    ws = wb.active
    ws.title = "Field Production"
    ws.append(list(field.columns))
    for row in field.itertuples(index=False):
        ws.append([row[0], row[1].to_pydatetime(), *[float(v) for v in row[2:]]])
    units = {"oil": "Oil production cumulative [STB]", "gas": "Gas production cumulative [MSCF]",
             "water": "Water production cumulative [STB]"}
    for k, title in (("gas", "Gas production cumulative"), ("oil", "Oil production cumulative"),
                     ("water", "Water production cumulative")):
        _write_curve_sheet(wb.create_sheet(title), stamps, sim[k], units[k])
    wb.save(out / "02 Field Production.xlsx")

    # ---- 09 Template
    f_dates = pd.date_range("2008-04-01", "2028-01-01", freq="QS-JAN")
    wb = Workbook()
    wb.remove(wb.active)
    for i in range(10):
        ws = wb.create_sheet(f"Case {i + 1}")
        ws.append(["Date", "Gas production cumulative [MSCF]", "Oil production cumulative [STB]",
                   "Water production cumulative [STB]"])
        for j, d in enumerate(f_dates):
            ws.append([d.to_pydatetime() if j == len(f_dates) - 1 else d.strftime("%m/%d/%Y")])
    wb.save(out / "09 Template Deliverable.xlsx")
    return {"truth": dict(zip(PARAMS, truth)), "params": params}
