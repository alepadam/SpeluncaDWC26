"""Grid workbooks (files 03-08): read case sheets and test how cases differ.

Findings on the real data (checked cell by cell on 35 cases, files 03 and 08, 15 Oct 2026): every case is the SAME
base grid with the SAME initial state, changed only by 3 of the 4 uncertainty multipliers:
  * porosity multiplier      -> PORO and PORV of every active cell
  * permeability multiplier  -> PERMX/PERMY (and so TRANX/TRANY) in one region of the lower layers; PERMZ unchanged
  * aquifer pore volume      -> PORV of the aquifer cells at the grid edge
  * fault transmissibility   -> NOT visible in any column (applied on fault faces inside the simulator)
All other columns are identical in every case. So the grids hold no information beyond the parameters; they are used
to verify this, to locate the oil zone / aquifer / permeability region, and to derive the initial hydrocarbon pore
volume (HCPV), which scales exactly with the porosity multiplier.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def column_codes(header: list) -> dict[str, int]:
    """'Pore Volume (PORV)' -> {'PORV': 6}. Matches on the code in brackets, robust to label typos/spaces."""
    out = {}
    for i, h in enumerate(header):
        s = str(h).strip() if h is not None else ""
        if s.endswith(")") and "(" in s:
            out[s[s.rindex("(") + 1:-1].strip()] = i
        elif s:
            out[s] = i
    return out


def read_grid_sheet(path: Path, sheet: str, null_value: float = -999.25) -> tuple[dict[str, int], np.ndarray]:
    """Whole sheet as a float array (nulls -> NaN) plus code -> column index. ~15 s per 90k-row sheet."""
    wb = load_workbook(path, read_only=True)
    it = wb[sheet].iter_rows(values_only=True)
    cols = column_codes(list(next(it)))
    arr = np.array([[np.nan if v is None else float(v) for v in row] for row in it])
    wb.close()
    arr[arr == null_value] = np.nan
    return cols, arr


def active_mask(cols: dict, arr: np.ndarray) -> np.ndarray:
    """Active cell = positive pore volume (inactive cells carry PORV 0 and -999.25 elsewhere)."""
    porv = arr[:, cols["PORV"]]
    return np.isfinite(porv) & (porv > 0)


def hcpv(cols: dict, arr: np.ndarray) -> float:
    """Initial hydrocarbon pore volume in reservoir barrels: sum of PORV x SOIL over active cells."""
    m = active_mask(cols, arr)
    return float(np.sum(arr[m, cols["PORV"]] * arr[m, cols["SOIL"]]))


def compare_cases(a: tuple, b: tuple, params_a: dict, params_b: dict) -> dict:
    """Check that case b = case a transformed by the multiplier ratios. Returns the measured vs expected ratios."""
    (ca, xa), (cb, xb) = a, b
    act = active_mask(ca, xa)
    out = {"same_cells": bool(np.array_equal(xa[:, :3], xb[:, :3])),
           "same_active_set": bool(np.array_equal(act, active_mask(cb, xb)))}
    r_poro = xb[act, cb["PORO"]] / xa[act, ca["PORO"]]
    out["poro_ratio"] = float(np.median(r_poro))
    out["poro_ratio_expected"] = params_b["Porosity Multiplier"] / params_a["Porosity Multiplier"]
    r_perm = xb[act, cb["PERMX"]] / xa[act, ca["PERMX"]]
    region = np.abs(r_perm - 1.0) > 1e-6
    out["perm_region_cells"] = int(region.sum())
    out["perm_ratio"] = float(np.median(r_perm[region])) if region.any() else 1.0
    out["perm_ratio_expected"] = params_b["Permeability Multiplier"] / params_a["Permeability Multiplier"]
    r_porv = xb[act, cb["PORV"]] / xa[act, ca["PORV"]]
    aq = np.abs(r_porv - out["poro_ratio"]) > 1e-3
    out["aquifer_cells"] = int(aq.sum())
    out["aquifer_ratio"] = float(np.median(r_porv[aq])) if aq.any() else out["poro_ratio"]
    out["aquifer_ratio_expected"] = (params_b["Aquifer Pore Volume"] / params_a["Aquifer Pore Volume"]
                                     * out["poro_ratio_expected"])
    out["initial_state_identical"] = all(
        np.allclose(xa[act, ca[c]], xb[act, cb[c]], equal_nan=True) for c in ("PRESSURE", "SWAT", "SOIL", "RS"))
    return out


# ------------------------------------------------------------------ whole-grid verification (notebook 01 section 9)
SCALED_COLUMNS = ("PORO", "PORV", "PERMX", "PERMY", "TRANX", "TRANY")


def ratios(p_base, p_case) -> dict:
    """Multiplier ratios case / base for the 4 parameters (inputs: rows of the Uncertainty Values table)."""
    return {"poro": p_case["Porosity Multiplier"] / p_base["Porosity Multiplier"],
            "perm": p_case["Permeability Multiplier"] / p_base["Permeability Multiplier"],
            "aquifer": p_case["Aquifer Pore Volume"] / p_base["Aquifer Pore Volume"],
            "fault": p_case["Fault Transmissibility"] / p_base["Fault Transmissibility"]}


def region_masks(base: tuple, other: tuple, poro_ratio: float) -> dict:
    """Cell groups found by comparing two cases: active cells, permeability region, aquifer cells, oil zone.

    Needs two cases whose permeability and aquifer multipliers differ (true for any two cases of the design)."""
    (c, a), (_, b) = base, other
    act = active_mask(c, a)
    with np.errstate(divide="ignore", invalid="ignore"):
        perm_region = act & (np.abs(b[:, c["PERMX"]] / a[:, c["PERMX"]] - 1) > 1e-3)
        aquifer = act & (np.abs(b[:, c["PORV"]] / a[:, c["PORV"]] / poro_ratio - 1) > 1e-3)
    return {"active": act, "perm_region": perm_region, "aquifer": aquifer, "oil_zone": act & (a[:, c["SOIL"]] > 0)}


def _max_rel_err(x: np.ndarray, pred: np.ndarray) -> float:
    """Largest |x / pred - 1| over cells whose predicted value is not tiny (tiny values carry export rounding)."""
    big = np.isfinite(pred) & (np.abs(pred) > 1e-3 * np.nanmax(np.abs(pred)))
    return float(np.nanmax(np.abs(x[big] / pred[big] - 1))) if big.any() else 0.0


def check_case(base: tuple, case: tuple, masks: dict, r: dict) -> dict:
    """Predict every column of `case` as base x multiplier ratios `r`; return the largest errors.

    Expected if the grid adds nothing: errors ~1e-5 or less (export rounding), unchanged columns identical,
    every TRANX/TRANY ratio between 1 and the permeability ratio, and no cell scaled by the fault ratio."""
    (cb, a), (cc, x) = base, case
    act, aq, reg = masks["active"], masks["aquifer"], masks["perm_region"]
    out = {"same_active_cells": bool(np.array_equal(act, active_mask(cc, x)))}
    pk = np.where(reg, r["perm"], 1.0)
    pred = {"PORO": a[:, cb["PORO"]] * r["poro"],
            "PORV": a[:, cb["PORV"]] * r["poro"] * np.where(aq, r["aquifer"], 1.0),
            "PERMX": a[:, cb["PERMX"]] * pk, "PERMY": a[:, cb["PERMY"]] * pk}
    for col, p in pred.items():
        out[f"{col}_max_rel_err"] = _max_rel_err(x[act, cc[col]], p[act])
    same = [k for k in cb if k not in SCALED_COLUMNS]
    out["unchanged_cols_max_abs_diff"] = float(max(
        np.nanmax(np.abs(x[act, cc[k]] - a[act, cb[k]])) if np.isfinite(a[act, cb[k]]).any() else 0.0 for k in same))
    lo, hi = min(1.0, r["perm"]), max(1.0, r["perm"])
    fault_testable = not (lo - 1e-3 <= r["fault"] <= hi + 1e-3)
    unexplained, fault_cells = 0, 0
    for col in ("TRANX", "TRANY"):
        t0, t1 = a[act, cb[col]], x[act, cc[col]]
        big = t0 > 1e-3 * np.nanmax(t0)
        q = t1[big] / t0[big]
        unexplained += int(np.sum((q < lo * (1 - 1e-3)) | (q > hi * (1 + 1e-3))))
        if fault_testable:
            fault_cells += int(np.sum(np.abs(q - r["fault"]) < 1e-3 * r["fault"]))
    out["tran_cells_unexplained"] = unexplained
    out["fault_testable"] = fault_testable
    out["fault_scaled_cells"] = fault_cells if fault_testable else None
    out["hcpv_rb"] = hcpv(cc, x)
    return out


def layout_summary(base: tuple, masks: dict) -> dict:
    """Where the oil zone, aquifer cells and permeability region sit, and their share of pore volume."""
    c, a = base
    act, pv = masks["active"], a[:, c["PORV"]]
    total = float(np.nansum(pv[act]))
    find = {code: next((k for k in c if k.upper() == code or k.upper().startswith(code + " ")), None)
            for code in ("I", "J", "K", "Z")}
    out = {}
    for name in ("oil_zone", "aquifer", "perm_region"):
        m = masks[name]
        out[name] = {"cells": int(m.sum()), "pore_volume_share": float(np.nansum(pv[m]) / total)}
        if find["Z"] is not None:
            out[name]["z_min"] = float(np.nanmin(a[m, c[find["Z"]]]))
            out[name]["z_max"] = float(np.nanmax(a[m, c[find["Z"]]]))
        for code in ("I", "J", "K"):
            if find[code] is not None:
                out[name][f"{code.lower()}_range"] = (int(np.nanmin(a[m, c[find[code]]])),
                                                      int(np.nanmax(a[m, c[find[code]]])))
    out["perm_region"]["in_oil_zone"] = int((masks["perm_region"] & masks["oil_zone"]).sum())
    out["aquifer"]["in_oil_zone"] = int((masks["aquifer"] & masks["oil_zone"]).sum())
    return out
