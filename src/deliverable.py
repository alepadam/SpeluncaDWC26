"""Fill '09 Template Deliverable.xlsx' with the best 10 forecasts (sheet 'Case i' = i-th best realization).

The template is copied, never edited in place. Dates in column A are written as real dates (the template
mixes text and date cells), values go under the Gas / Oil / Water headers found in row 1.
"""

from __future__ import annotations

import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .config import find_raw_file
from .io_load import load_template_layout
from .utils import get_logger

HEADER_KEYS = {"gas": "gas production cumulative", "oil": "oil production cumulative",
               "water": "water production cumulative"}


def run(cfg: dict, state: dict) -> dict:
    log = get_logger("deliverable", cfg["paths"]["logs"])
    src = find_raw_file(cfg, cfg["files"]["template"])
    dst = cfg["paths"]["deliverable"] / cfg["deliverable"]["output_name"]
    layout = load_template_layout(cfg)
    f_dates = pd.DatetimeIndex(state["forecast_dates"])
    fc = state["best10_forecasts"]
    n = fc["oil"].shape[0]
    expected = [f"Case {i + 1}" for i in range(n)]
    if list(layout) != expected:
        raise ValueError(f"Template sheets {list(layout)} do not match {expected}")
    for sheet, dates in layout.items():
        if not pd.DatetimeIndex(dates).equals(f_dates):
            raise ValueError(f"Sheet {sheet!r}: template dates differ from the forecast dates")

    for k in HEADER_KEYS:                       # cumulatives must never decrease (F46)
        bad = np.flatnonzero((np.diff(fc[k], axis=1) < -1e-6).any(axis=1))
        if bad.size:
            raise ValueError(f"Cumulative {k} decreases in {[expected[j] for j in bad]}")
    shutil.copyfile(src, dst)
    wb = load_workbook(dst)
    for i, sheet in enumerate(expected):
        ws = wb[sheet]
        header = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}
        cols = {}
        for k, key in HEADER_KEYS.items():
            hits = [col for h, col in header.items() if h.startswith(key)]
            if len(hits) != 1:
                raise KeyError(f"Sheet {sheet!r}: expected one column starting with {key!r}, found {hits}")
            cols[k] = hits[0]
        for r, d in enumerate(f_dates, start=2):
            cell = ws.cell(row=r, column=1, value=d.to_pydatetime())
            cell.number_format = cfg["deliverable"]["date_format"]
            for k, col in cols.items():
                v = float(fc[k][i, r - 2])
                if not np.isfinite(v) or v < 0:
                    raise ValueError(f"{sheet} {d.date()} {k}: invalid value {v}")
                ws.cell(row=r, column=col, value=v)
    wb.save(dst)
    wb.close()

    # verify by reading back
    check = load_workbook(dst, read_only=True, data_only=True)
    for i, sheet in enumerate(expected):
        rows = list(check[sheet].iter_rows(min_row=2, max_row=len(f_dates) + 1, values_only=True))
        if len(rows) != len(f_dates) or any(v is None for row in rows for v in row[:4]):
            raise ValueError(f"Read-back check failed on {sheet!r}")
    check.close()
    log.info("Deliverable written: %s (%d sheets x %d quarters)", dst, n, len(f_dates))
    return {"deliverable_path": dst}
