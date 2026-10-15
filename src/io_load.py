"""Faithful readers for every DWC26 input file. No modelling logic here."""

from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from .config import find_raw_file
from .utils import norm_case


# ------------------------------------------------------------------ 01 Introduction
def load_uncertainty(cfg: dict) -> pd.DataFrame:
    """100 cases x 4 parameters, indexed by normalised case id ('CASE 1')."""
    path = find_raw_file(cfg, cfg["files"]["introduction"])
    df = pd.read_excel(path, sheet_name=cfg["sheets"]["uncertainty"])
    df.columns = [str(c).strip() for c in df.columns]
    names = cfg["parameters"]["names"]
    missing = [c for c in [cfg["parameters"]["case_col"], *names] if c not in df.columns]
    if missing:
        raise KeyError(f"Uncertainty sheet is missing columns {missing}; found {list(df.columns)}")
    df = df.dropna(subset=[cfg["parameters"]["case_col"]])
    df["case_id"] = df[cfg["parameters"]["case_col"]].map(norm_case)
    df = df.set_index("case_id")[names].astype(float)
    if df.index.duplicated().any():
        raise ValueError(f"Duplicate case ids: {df.index[df.index.duplicated()].tolist()}")
    return df


# ------------------------------------------------------------------ 02 Field Production
def list_field_sheets(cfg: dict) -> list[str]:
    path = find_raw_file(cfg, cfg["files"]["field_production"])
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def load_observed(cfg: dict) -> pd.DataFrame:
    """Observed field rates indexed by Date; columns renamed to oil, gas, water, winj (per day)."""
    oc = cfg["observed_columns"]
    path = find_raw_file(cfg, cfg["files"]["field_production"])
    df = pd.read_excel(path, sheet_name=cfg["sheets"]["field"])
    df.columns = [str(c).strip() for c in df.columns]
    rename = {oc["oil_rate"]: "oil", oc["gas_rate"]: "gas", oc["water_rate"]: "water", oc["water_inj"]: "winj"}
    missing = [c for c in [oc["date"], *rename] if c not in df.columns]
    if missing:
        raise KeyError(f"Field Production sheet is missing columns {missing}; found {list(df.columns)}")
    df[oc["date"]] = pd.to_datetime(df[oc["date"]])
    df = df.rename(columns=rename).set_index(oc["date"]).sort_index()
    df.index.name = "Date"
    return df[["oil", "gas", "water", "winj"]].astype(float)


def load_curve_sheet(cfg: dict, sheet: str) -> tuple[pd.DataFrame, pd.Series | None]:
    """One simulation-curve sheet -> (cases DataFrame, RAMP_History Series or None).

    Layout: row 0 = Y1..Y101, row 1 = Case_1..Case_100 + RAMP_History, row 2 = units,
    data from row 3; column 0 = irregular timestamps. Blank cells stay NaN.
    """
    path = find_raw_file(cfg, cfg["files"]["field_production"])
    raw = pd.read_excel(path, sheet_name=sheet, header=None)
    labels = [str(x).strip() for x in raw.iloc[cfg["sheets"]["case_label_row"], 1:].tolist()]
    data = raw.iloc[cfg["sheets"]["data_start_row"]:].copy()
    data.columns = ["Date", *labels]
    data = data.dropna(subset=["Date"])
    data["Date"] = pd.to_datetime(data["Date"])
    data = data.set_index("Date").sort_index().apply(pd.to_numeric, errors="coerce")
    data = data.loc[:, [c for c in data.columns if c and c.lower() != "nan"]]

    hist_name = cfg["sheets"]["history_column"]
    history = data.pop(hist_name) if hist_name in data.columns else None
    data.columns = [norm_case(c) for c in data.columns]
    bad = [c for c in data.columns if not c.startswith("CASE ")]
    if bad:
        raise ValueError(f"Unexpected columns in sheet {sheet!r}: {bad}")
    if data.columns.duplicated().any():
        raise ValueError(f"Duplicate case columns in sheet {sheet!r}")
    return data, history


# ------------------------------------------------------------------ 09 Template
def load_template_layout(cfg: dict) -> dict[str, list[pd.Timestamp]]:
    """Sheet name -> forecast dates in column A. Dates are text 'MM/DD/YYYY' except possibly the last."""
    path = find_raw_file(cfg, cfg["files"]["template"])
    wb = load_workbook(path, read_only=True)
    layout = {}
    for ws in wb.worksheets:
        dates = []
        for (val,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if val is None or (isinstance(val, str) and not val.strip()):
                continue
            dates.append(pd.to_datetime(val.strip(), format="%m/%d/%Y") if isinstance(val, str) else pd.Timestamp(val))
        layout[ws.title] = dates
    wb.close()
    return layout

